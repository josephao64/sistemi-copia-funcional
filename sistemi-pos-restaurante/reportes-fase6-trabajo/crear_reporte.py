from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root = Path(r'C:\Users\Elizabeth\Downloads\sistema rest\sistemi-pos-restaurante')
source = root / 'docs/FASE_6_REPORTE_Y_MANUAL_DE_USUARIO.md'
destination = root / 'docs/FASE_6_REPORTE_Y_MANUAL_DE_USUARIO.docx'
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Cm(21), Cm(29.7)
section.top_margin, section.bottom_margin = Cm(1.75), Cm(1.65)
section.left_margin = section.right_margin = Cm(1.8)
section.header_distance = section.footer_distance = Cm(.7)
styles = doc.styles
for border in list(styles.element.iter(qn('w:pBdr'))):
    border.getparent().remove(border)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3', 'List Paragraph']:
    style = styles[name]
    style.font.name = 'Arial'
    style.font.color.rgb = RGBColor(0, 0, 0)
    for child in list(style.element):
        if child.tag == qn('w:rPr'):
            for color in child.findall(qn('w:color')):
                color.attrib.pop(qn('w:themeColor'), None)
normal = styles['Normal']
normal.font.size = Pt(10.5)
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.line_spacing = 1.10
normal.paragraph_format.widow_control = True
styles['Title'].font.size = Pt(25)
styles['Title'].font.bold = True
styles['Title'].paragraph_format.space_after = Pt(16)
for name, size, before, after in [('Heading 1', 18, 0, 12), ('Heading 2', 12.5, 12, 6), ('Heading 3', 11, 8, 5)]:
    s = styles[name]
    s.font.size = Pt(size); s.font.bold = True
    s.paragraph_format.space_before = Pt(before)
    s.paragraph_format.space_after = Pt(after)
    s.paragraph_format.keep_with_next = True
    s.paragraph_format.keep_together = True
styles['List Paragraph'].paragraph_format.space_after = Pt(5)
styles['List Paragraph'].paragraph_format.left_indent = Cm(.4)
styles['List Paragraph'].paragraph_format.first_line_indent = Cm(-.4)

header = section.header.paragraphs[0]
header.text = 'SISTEMI POS Restaurante  |  Fase 6'
header.runs[0].font.name = 'Arial'; header.runs[0].font.size = Pt(8)
header.runs[0].font.color.rgb = RGBColor(0,0,0)
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('Manual de usuario  ·  Página ')
for field in ['PAGE', 'NUMPAGES']:
    if field == 'NUMPAGES': footer.add_run(' de ')
    begin = OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'), 'begin')
    instruction = OxmlElement('w:instrText'); instruction.set(qn('xml:space'), 'preserve'); instruction.text = f' {field} '
    end = OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'), 'end')
    footer.add_run()._r.append(begin); footer.add_run()._r.append(instruction); footer.add_run()._r.append(end)
for run in footer.runs: run.font.size = Pt(8); run.font.name = 'Arial'

def inline(paragraph, text):
    for part in re.split(r'(\*\*.*?\*\*|`[^`]+`)', text):
        if not part: continue
        run = paragraph.add_run(part[2:-2] if part.startswith('**') else part[1:-1] if part.startswith('`') else part)
        if part.startswith('**'): run.bold = True
        if part.startswith('`'): run.font.name = 'Consolas'; run.font.size = Pt(9)

def table(rows):
    columns = len(rows[0])
    if columns == 2: widths = [2.1, 4.75]
    elif columns == 4: widths = [3.55, 1.1, 1.1, 1.1]
    elif rows[0][0] == 'Clave': widths = [.7, 1.65, 4.5]
    elif rows[0][0] == 'Estado visible': widths = [1.15, 3.25, 2.45]
    elif rows[0][0] == 'Capa o componente': widths = [1.2, 2.75, 2.9]
    elif rows[0][0] == 'Sección': widths = [1.0, 3.3, 2.55]
    else: widths = [2.05, 2.2, 2.6]
    t = doc.add_table(rows=0, cols=columns)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    for col, w in zip(t.columns, widths): col.width = Inches(w)
    props = t._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        item = OxmlElement('w:'+edge)
        for k,v in [('val','single'),('sz','4'),('color','D9D9D9')]: item.set(qn('w:'+k), v)
        borders.append(item)
    props.append(borders)
    for i, row in enumerate(rows):
        cells = t.add_row().cells
        trpr = cells[0]._tc.getparent().get_or_add_trPr()
        trpr.append(OxmlElement('w:cantSplit'))
        if i == 0: trpr.append(OxmlElement('w:tblHeader'))
        for j, (cell, value, width) in enumerate(zip(cells,row,widths)):
            cell.width = Inches(width); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), '203E59' if i == 0 else ('F3F6F8' if i % 2 == 0 else 'FFFFFF')); tcpr.append(shade)
            margins = OxmlElement('w:tcMar')
            for edge, val in [('top','90'),('bottom','90'),('left','105'),('right','105')]:
                item = OxmlElement('w:'+edge); item.set(qn('w:w'),val); item.set(qn('w:type'),'dxa'); margins.append(item)
            tcpr.append(margins)
            p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.05
            if i == 0 or (columns == 4 and j > 0) or (rows[0][0] == 'Clave' and j == 0): p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            inline(p,value)
            for run in p.runs:
                run.font.size = Pt(9.5); run.font.name = 'Arial'
                if i == 0: run.bold = True; run.font.color.rgb = RGBColor(255,255,255)
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(3); p.paragraph_format.space_before = Pt(0); p.paragraph_format.line_spacing = 1
    p.add_run().font.size = Pt(2)

lines = source.read_text(encoding='utf-8').splitlines()
i = 0
while i < len(lines):
    line = lines[i].strip()
    if not line: i += 1; continue
    if line == '<!-- pagebreak -->': doc.add_page_break(); i += 1; continue
    if line.startswith('|'):
        rows = []
        while i < len(lines) and lines[i].strip().startswith('|'):
            values = [v.strip() for v in lines[i].strip().strip('|').split('|')]
            if not all(re.fullmatch(r'[:\- ]+', v) for v in values): rows.append(values)
            i += 1
        table(rows); continue
    if line.startswith('# '): inline(doc.add_paragraph(style='Title'), line[2:])
    elif line.startswith('## '): inline(doc.add_paragraph(style='Heading 1'), line[3:])
    elif line.startswith('### '): inline(doc.add_paragraph(style='Heading 2'), line[4:])
    elif re.match(r'^(\d+\.|-) ', line): inline(doc.add_paragraph(style='List Paragraph'), '• '+line[2:] if line.startswith('- ') else line)
    else: inline(doc.add_paragraph(),line)
    i += 1

doc.core_properties.title = 'Reporte y manual de usuario de la fase 6'
doc.core_properties.subject = 'Pedidos cocina y delivery de SISTEMI POS Restaurante'
doc.core_properties.author = 'Documentación del proyecto SISTEMI POS'
doc.core_properties.keywords = 'Fase 6, pedidos, cocina, delivery, manual'
doc.core_properties.comments = 'Basado en la implementación local del 19 de octubre de 2026.'
doc.save(destination)
print(destination)
print(f'{len(doc.paragraphs)} párrafos y {len(doc.tables)} tablas')
