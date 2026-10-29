from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

out = Path(r'C:\Users\Elizabeth\Downloads\sistema rest\sistemi-pos-restaurante\docs\entrega_fase6')
out.mkdir(parents=True, exist_ok=True)
doc = Document()
for border in list(doc.styles.element.iter(qn('w:pBdr'))):
    border.getparent().remove(border)
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin, sec.bottom_margin = Inches(.72), Inches(.68)
sec.left_margin = sec.right_margin = Inches(.8)
sec.header_distance = sec.footer_distance = Inches(.3)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3', 'List Paragraph']:
    style = doc.styles[name]
    style.font.name = 'Arial'
    style.font.color.rgb = RGBColor(0,0,0)
    style.paragraph_format.widow_control = True
normal=doc.styles['Normal']; normal.font.size=Pt(10.5)
normal.paragraph_format.line_spacing=1.12
normal.paragraph_format.space_after=Pt(7)
doc.styles['Title'].font.size=Pt(27); doc.styles['Title'].font.bold=True
doc.styles['Title'].paragraph_format.space_after=Pt(12)
doc.styles['Subtitle'].font.size=Pt(14)
for name,size in [('Heading 1',18),('Heading 2',12),('Heading 3',11)]:
    style=doc.styles[name]; style.font.size=Pt(size); style.font.bold=True
    style.paragraph_format.space_before=Pt(12 if name!='Heading 1' else 0)
    style.paragraph_format.space_after=Pt(7)
    style.paragraph_format.keep_with_next=True
header=sec.header.paragraphs[0]
header.text='SISTEMI POS Restaurante   |   Informe de entrega técnica'
for run in header.runs: run.font.size=Pt(8)
footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('Fase 6   ·   20 de octubre de 2026   ·   Página ')
field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); footer._p.append(field)
for run in footer.runs: run.font.size=Pt(8)

def p(text,bold=False):
    par=doc.add_paragraph(); run=par.add_run(text); run.bold=bold; return par
def h(text,level=2): doc.add_heading(text,level)
def lead(label,text):
    par=doc.add_paragraph(); par.add_run(label+' ').bold=True; par.add_run(text); return par
def page(title): doc.add_page_break(); h(title,1)
def table(headers,rows,widths):
    t=doc.add_table(rows=1,cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    for col,width in zip(t.columns,widths): col.width=Inches(width)
    for i,title in enumerate(headers): t.rows[0].cells[i].text=title
    repeat=OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for values in rows:
        cells=t.add_row().cells
        for cell,text in zip(cells,values): cell.text=text
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+edge); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'D9D9D9'); borders.append(e)
    t._tbl.tblPr.append(borders)
    for ridx,row in enumerate(t.rows):
        no_split=OxmlElement('w:cantSplit'); row._tr.get_or_add_trPr().append(no_split)
        for idx,cell in enumerate(row.cells):
            cell.width=Inches(widths[idx]); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props=cell._tc.get_or_add_tcPr(); margins=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                m=OxmlElement('w:'+edge); m.set(qn('w:w'),'90'); m.set(qn('w:type'),'dxa'); margins.append(m)
            props.append(margins)
            shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'17324D' if ridx==0 else ('F3F6F8' if ridx%2==0 else 'FFFFFF')); props.append(shade)
            for par in cell.paragraphs:
                par.paragraph_format.space_after=Pt(1); par.paragraph_format.line_spacing=1.04
                for run in par.runs:
                    run.font.name='Arial'; run.font.size=Pt(9.4)
                    if ridx==0: run.bold=True; run.font.color.rgb=RGBColor(255,255,255)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)
    return t

doc.add_paragraph('Informe de entrega de la fase 6',style='Title')
doc.add_paragraph('Pedidos cocina y delivery',style='Subtitle')
p('SISTEMI POS Restaurante',True)
p('Fecha de corte 20 de octubre de 2026')
h('Resultado de la entrega')
p('El proyecto incorpora el desarrollo del alcance principal ORD-01 a ORD-10: captura y revisión de pedidos, preparación por estación, mesas, entregas y control de reservas, consumos y cancelaciones. La entrega incluye el rediseño de las pantallas y la configuración del menú por sucursal.')
p('El módulo permite seguir un pedido desde su borrador hasta la entrega y conserva la información comercial y de inventario necesaria para enlazarlo posteriormente con ventas y pagos. Su implementación reutiliza catálogo, recetas, bodegas, lotes, permisos y movimientos de las fases anteriores.')
table(['Aspecto','Estado documentado'],[
('Alcance principal','Diez funcionalidades con implementación identificada en el código.'),
('Verificación técnica','Última compilación registrada: 0 errores y 0 advertencias. JavaScript sin errores de sintaxis en los archivos revisados.'),
('Aceptación de operación','Pendiente de evidencia funcional, concurrencia, migración ensayada y conciliación.'),
('Extras avanzados','Interfaz incorporada; algunas integraciones requieren corrección o verificación antes del cierre definitivo.')
],[1.55,5.35])
h('Propósito del informe')
p('Presentar lo entregado respecto de la tabla de fase 6 proporcionada por el usuario, explicar los extras incorporados y dejar un registro verificable del estado de cierre. El informe acredita la entrega de desarrollo; la aceptación operativa debe completarse con los escenarios descritos al final.')

page('Cobertura del alcance original')
p('Cada fila relaciona la clave del plan con el resultado incorporado. El estado corresponde a revisión de código y documentación; no sustituye la ejecución del criterio de aceptación en una base de prueba.')
table(['Clave','Entregable','Resultado incorporado'],[
('ORD-01','Canales y tipos de servicio','Canal independiente de salón, llevar o delivery. Sucursal, bodega y referencia externa forman parte del pedido histórico.'),
('ORD-02','Pedido y comanda','Borradores, líneas, revisión, numeración por secuencia y estados hasta entrega o cancelación. Las transiciones se controlan en el servidor.'),
('ORD-03','Menú por sucursal','Carta con secciones propias, fotos, favoritos, precio local y disponibilidad base. Selector buscable de artículos vendibles y estaciones.'),
('ORD-04','Modificadores y extras','Grupos con mínimos y máximos; opciones con precio y componentes adicionales. El pedido enviado conserva la configuración calculada.'),
('ORD-05','Exclusiones','Componentes directos permitidos de la receta activa elaborada al pedido. La exclusión ajusta consumo y conserva precio salvo regla expresa.'),
('ORD-06','Combos y promociones','Paquetes con elecciones, vigencias y promociones por prioridad. Prorrateo del precio y registro del consumo por componente.'),
('ORD-07','Mesas y salones','Áreas, mesas y turnos activos. Apertura, cambio controlado a mesa libre y liberación cuando los pedidos terminan operativamente.'),
('ORD-08','KDS de cocina','Cola por estación y canal; preparación total o parcial, marcado de listo, tiempos y acceso al historial de cambios.'),
('ORD-09','Delivery operativo','Cliente, teléfono, dirección, zona y tarifa. Motorista, ruta, salida, incidencias y confirmación de entrega.'),
('ORD-10','Cancelación y desperdicio','Reserva al enviar, consumo al iniciar preparación, liberación de lo pendiente, merma de lo preparado y retorno intacto autorizado.')
],[.65,1.55,4.7])
p('El objetivo del plan se conserva: pedido, preparación, entrega y venta son procesos relacionados con estados y efectos propios. El cobro y el cierre financiero forman parte de la integración de fase 7.')

page('Extras incorporados al alcance')
h('Menú y configuración de la sucursal')
lead('Secciones de carta.', 'Grupos locales como Pizzas, Acompañamientos y Bebidas, independientes de las categorías de inventario. Los productos anteriores sin sección conservan sus registros hasta asignarlos.')
lead('Selección de productos.', 'Búsqueda y filtros del catálogo vendible. Los insumos, ingredientes, subrecetas y empaques se excluyen de la carta; los preparados necesitan receta activa.')
lead('Recetas y exclusiones.', 'Detección de la receta al elegir el producto y selección de sus componentes directos cuando se elabora al pedir. El guardado en el servidor comprueba que las exclusiones correspondan a la receta activa.')
lead('Editor comercial.', 'Interfaz con tres columnas para secciones, productos e inspector. Incluye opciones, extras, estación de preparación, favoritos, orden y vista previa del menú.')
h('Experiencia de uso y operación')
lead('Diseño consistente.', 'Colores del sistema, fondos neutros, bordes discretos, controles uniformes y fotografías o iniciales en los artículos sin imagen. Se conserva la ubicación de los paneles.')
lead('Consulta y seguimiento.', 'Indicadores de pedidos, filtros, mensajes sin resultados, cocina por estado, mesas agrupadas por área y tarjetas de delivery.')
lead('Acciones guiadas.', 'Diálogos para confirmar operaciones, indicar cantidades y registrar motivos. La cocina evita actualizarse automáticamente mientras hay un diálogo o una operación pendiente.')
lead('Preparación parcial.', 'División de una línea en cantidades preparadas y pendientes, manteniendo importes y reservas sin duplicarlos.')
lead('Resguardo operativo.', 'Revisión de pedido y configuración, identificadores de reintento, referencia externa única, permisos por acción y restricciones de sucursal en el servidor.')
h('Extras con integración pendiente')
p('La interfaz también incorpora copia de menú entre sucursales, nombre y precio comercial personalizados, canales y horarios por producto y simulación del pedido. No se presentan como funciones aceptadas de extremo a extremo: la revisión detectó pendientes concretos que se detallan en la página final.')

page('Flujo operativo y control de inventario')
h('Recorrido del pedido')
lead('1 Captura.', 'El operador elige sucursal y bodega de consumo, canal, servicio, cliente y productos. El borrador calcula importes y admite revisión sin reservar inventario.')
lead('2 Envío.', 'El servidor valida recetas, opciones, precio y existencia. Guarda los datos históricos y reserva artículos o componentes por bodega y lote.')
lead('3 Preparación.', 'Cocina inicia toda la línea o una parte. La reserva pasa a consumo físico una sola vez. Las cantidades que siguen pendientes permanecen reservadas.')
lead('4 Despacho y entrega.', 'Se marca listo cuando corresponde. Delivery agrega asignación de motorista y salida a ruta antes de confirmar la recepción del cliente.')
lead('5 Seguimiento comercial.', 'El pedido y sus consumos quedan disponibles para el contrato de integración de ventas. Confirmar entrega o liberar mesa no registra un pago.')
h('Ejemplo de control físico')
p('Si una bodega tiene 10 bebidas y un pedido reserva 3, conserva 10 unidades físicas, 3 reservadas y 7 disponibles. Al iniciar el despacho de esas 3, registra una salida y queda con 7 físicas y ninguna reserva de esa línea. Marcar listo o entregar no genera otra salida.')
table(['Situación','Tratamiento incorporado'],[
('Cancelar borrador','No hubo reserva ni consumo; se cancela el importe activo.'),
('Cancelar cantidad pendiente','Libera su reserva; conserva la existencia física.'),
('Cancelar cantidad preparada','Conserva el consumo y registra la merma con motivo y trazabilidad.'),
('Reintegrar artículo intacto','Requiere permisos, retorno físico elegible y referencia al movimiento y lote originales.')
],[2.05,4.85])
p('Para una pizza que utiliza masa ya producida, la preparación consume la masa terminada y los demás componentes de su receta. La harina utilizada al producir la masa no debe descontarse otra vez al atender el pedido.')

page('Estructura técnica y trazabilidad')
h('Integración con las fases anteriores')
p('La estructura empresarial aporta empresa, sucursal, bodega y alcance del usuario. El catálogo aporta tipos, variantes, unidades y configuración local. Inventario y compras aportan existencias y lotes; producción y recetas aportan versiones, componentes, rendimientos y costos. Fase 6 coordina estas bases durante la atención del pedido.')
table(['Capa','Componentes principales'],[
('Dominio','RestaurantOrder y RestaurantSetup: pedido, líneas, estados, configuración, reservas, eventos, merma y turnos.'),
('Aplicación','RestaurantService, RestaurantPricing y RestaurantPreparation: alcance, precios, recetas, preparación, cancelación y delivery.'),
('Persistencia','RestaurantRepository, RestaurantConfiguration y AppDbContext: consultas, transacciones, restricciones e históricos.'),
('Interfaz','RestaurantController, siete vistas Restaurant, restaurant.js, restaurant-config.js, restaurant-ui.js y restaurant.css.'),
('Inventario compartido','InventoryStock, InventoryMovement e InventoryLotLedger: reservas, consumos y protección de cantidades comprometidas.')
],[1.35,5.55])
h('Controles incorporados')
p('Las operaciones críticas comprueban permiso y sucursal en el servidor. Los pedidos y configuraciones usan revisiones para detectar cambios de otro operador. La persistencia utiliza transacciones serializables, secuencias, restricciones únicas y control de concurrencia PostgreSQL.')
p('Cada pedido conserva eventos, usuario, fecha y cambios. Las líneas enviadas guardan su configuración comercial y consumo previsto; los movimientos físicos y lotes permiten rastrear preparación, cancelación, merma y retorno elegible.')
h('Migración y contrato de salida')
p('La migración AddRestaurantOrdersKitchenDelivery incorpora las tablas operativas del módulo, la secuencia de pedidos y campos para reservas y movimientos relacionados. La aplicación de la migración sobre históricos y su conciliación no cuentan con evidencia de aceptación en esta entrega.')
p('El endpoint de checkout entrega identidad y revisión del pedido, ubicación, importes, líneas y movimientos de consumo existentes. Fase 7 debe utilizar esa información para crear la venta y el cobro sin repetir una salida física ya registrada.')

page('Evidencia de cierre y aceptación pendiente')
h('Verificación técnica disponible')
p('La última compilación documentada de Web y sus dependencias terminó con 0 errores y 0 advertencias. Los archivos JavaScript revisados pasaron su comprobación de sintaxis; git diff --check no reportó errores de formato. Son resultados de entregas previas de esta conversación, no una suite funcional ejecutada para este informe.')
p('Para preparar este documento se revisaron el código local, la guía de implementación, el manual y el registro del rediseño. No se ejecutaron pruebas de interfaz, concurrencia o migración de la base activa. La captura proporcionada evidencia una pantalla de configuración, pero no valida los diez recorridos operativos.')
h('Observaciones concretas de la revisión')
lead('Copia entre sucursales.', 'La copia genera identificadores nuevos para secciones y estaciones, mientras conserva las referencias anteriores en las líneas del menú. Debe remapear sus relaciones y validarse el guardado antes de declararse completada.')
lead('Precio y nombre personalizados.', 'El menú de consulta utiliza valores personalizados, pero la cotización que construye las líneas del pedido conserva el nombre y el cálculo de precio del catálogo. Debe unificarse el cálculo y comprobarse que pantalla, borrador y envío coinciden.')
lead('Disponibilidad y canales.', 'Los campos de canales y horarios existen en la configuración; no se identificó su aplicación completa al cotizar o enviar. La visibilidad de la interfaz debe estar respaldada por la misma validación en el servidor.')
lead('Exclusiones y reglas obligatorias.', 'La interfaz identifica algunos componentes como requeridos por su nombre. Esa indicación debe revisarse antes de presentarla como una regla explícita y garantizada por el servidor.')
h('Criterio para el cierre operativo')
p('Completar una migración ensayada y conciliada, el recorrido de salón, llevar y delivery, preparación parcial, cancelación pendiente y preparada, reintentos, permisos y acceso cruzado entre sucursales. Probar dos operadores sobre el último stock, un pedido y una mesa, y registrar los resultados. Las listas actuales cargan el historial de la sucursal; la capacidad con volumen sigue pendiente de medición.')
h('Conclusión de entrega')
p('Se entrega el desarrollo del núcleo de pedidos, cocina y delivery con cobertura identificada para ORD-01 a ORD-10 y mejoras de configuración e interfaz. El estado que respalda la evidencia es entrega técnica con aceptación operativa pendiente. El cierre definitivo requiere resolver las observaciones de extras y documentar la aceptación funcional; la integración de venta y cobro se conserva en fase 7.')

doc.core_properties.title='Informe de entrega de la fase 6'
doc.core_properties.subject='Pedidos cocina y delivery con cobertura y estado de aceptación'
doc.core_properties.author='SISTEMI POS'
doc.save(out/'Informe_Entrega_Fase_6_SISTEMI_POS.docx')
print(str(out/'Informe_Entrega_Fase_6_SISTEMI_POS.docx'))
